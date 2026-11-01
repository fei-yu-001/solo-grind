# -*- coding: utf-8 -*-
"""把简单的 Markdown 转成规范排版的 docx（标题黑体、正文宋体、三线表、代码块）。"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement



CJK = r'一-龥　-〿！-｠'
_TL = re.compile(r'(?<=[' + CJK + r'])[ 	　]+(?=[A-Za-z0-9/@#])')
_TR = re.compile(r'(?<=[A-Za-z0-9/@#])[ 	　]+(?=[' + CJK + r'])')


def tight(text):
    if not text:
        return text
    text = str(text).replace('`', '')
    m = re.match(r'^(\d+(?:\.\d+)*\.)[ 	]+', text)
    prefix = m.group(0) if m else ''
    body = text[len(prefix):]
    return prefix + _TR.sub('', _TL.sub('', body))


def set_run(run, cn='宋体', en='Times New Roman', size=12, bold=False):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:ascii'), en)
    rf.set(qn('w:hAnsi'), en)
    rf.set(qn('w:eastAsia'), cn)


def shade(p, fill='F5F5F5'):
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    ppr.append(shd)


def add_table(doc, rows):
    n = len(rows[0])
    widths = []
    for ci in range(n):
        longest = max(len(str(r[ci])) for r in rows)
        widths.append(max(2.0, min(8.0, longest * 0.42)))
    total = sum(widths)
    widths = [w / total * 14.0 for w in widths]
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblpr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'bottom'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '12')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    for edge in ('left', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'none')
        el.set(qn('w:sz'), '0')
        el.set(qn('w:color'), 'auto')
        borders.append(el)
    tblpr.append(borders)
    for ri, row in enumerate(rows):
        if ri > 0:
            t.add_row()
        for ci, val in enumerate(row):
            cell = t.rows[ri].cells[ci]
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(tight(val)), size=10)
            cell.width = Cm(widths[ci])
            if ri == 0:
                tcpr = cell._tc.get_or_add_tcPr()
                tb = OxmlElement('w:tcBorders')
                b = OxmlElement('w:bottom')
                b.set(qn('w:val'), 'single')
                b.set(qn('w:sz'), '6')
                b.set(qn('w:color'), '000000')
                tb.append(b)
                tcpr.append(tb)


def convert(src, dst):
    lines = open(src, encoding='utf-8').read().split('\n')
    doc = Document()
    for s in doc.sections:
        s.left_margin = Cm(2.8)
        s.right_margin = Cm(2.8)
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith('```'):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(4)
            pf.space_after = Pt(6)
            pf.line_spacing = 1.0
            pf.left_indent = Cm(0.4)
            r = p.add_run('\n'.join(buf))
            set_run(r, 'Consolas', 'Consolas', 9)
            shade(p)
            continue
        if line.startswith('|') and i + 1 < len(lines) and set(lines[i + 1].replace('|', '').strip()) <= set('-: '):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not (set(''.join(cells)) <= set('-: ')):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            doc.add_paragraph()
            continue
        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            level = len(m.group(1))
            sizes = {1: 18, 2: 15, 3: 13, 4: 12}
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14 if level <= 2 else 8)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.3
            set_run(p.add_run(tight(m.group(2))), '黑体', 'Arial', sizes[level], True)
            i += 1
            continue
        if line.startswith('---'):
            i += 1
            continue
        if re.match(r'^[-*]\s+', line) or re.match(r'^\d+\.\s+', line):
            text = re.sub(r'^[-*]\s+', '', line)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.5
            set_run(p.add_run(tight(text)), size=12)
            i += 1
            continue
        if line.startswith('>'):
            text = line.lstrip('> ').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.space_after = Pt(2)
            set_run(p.add_run(tight(text)), '楷体', 'Times New Roman', 10.5)
            i += 1
            continue
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', line)
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run(tight(text)), size=12)
        i += 1

    doc.save(dst)
    print('saved', dst)


if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])